import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# 1. Initialize 16:9 Widescreen Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# 2. Executive Design Palette (Modern Enterprise HR Slate / Indigo / Teal)
COLOR_BG_DARK     = RGBColor(15, 23, 42)      # Slate 900 #0F172A
COLOR_BG_LIGHT    = RGBColor(248, 250, 252)   # Slate 50 #F8FAFC
COLOR_PRIMARY     = RGBColor(15, 23, 42)      # Primary Text
COLOR_MUTED       = RGBColor(100, 116, 139)   # Slate 500
COLOR_ACCENT      = RGBColor(79, 70, 229)     # Indigo 600 #4F46E5
COLOR_TEAL        = RGBColor(13, 148, 136)    # Teal 600 #0D9488
COLOR_CARD_BG     = RGBColor(255, 255, 255)   # Pure White Card
COLOR_CARD_BORDER = RGBColor(226, 232, 240)   # Slate 200 Border
COLOR_GREEN       = RGBColor(16, 185, 129)    # Emerald 500
COLOR_AMBER       = RGBColor(245, 158, 11)    # Amber 500

def set_slide_background(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_header(slide, tag, title, dark=False):
    # Category Tag Pill
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.45), Inches(3.4), Inches(0.32))
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor(30, 41, 59) if dark else RGBColor(238, 242, 255)
    pill.line.color.rgb = RGBColor(99, 102, 241) if dark else RGBColor(199, 210, 254)
    tf = pill.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = tag.upper()
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = RGBColor(129, 140, 248) if dark else COLOR_ACCENT
    p.alignment = PP_ALIGN.CENTER
    
    # Main Slide Title
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.75))
    tf2 = tb.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = title
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(255, 255, 255) if dark else COLOR_PRIMARY

def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card

# ==============================================================================
# SLIDE 1: Cover Slide
# ==============================================================================
s1 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s1, COLOR_BG_DARK)

c1 = add_card(s1, Inches(0.8), Inches(1.1), Inches(11.733), Inches(5.3), RGBColor(30, 41, 59), RGBColor(51, 65, 85))
tf1 = c1.text_frame
tf1.word_wrap = True
tf1.margin_left = Inches(0.8)
tf1.margin_top = Inches(0.7)

p = tf1.paragraphs[0]
p.text = "DAIKO PEOPLE"
p.font.size = Pt(46)
p.font.bold = True
p.font.color.rgb = RGBColor(255, 255, 255)

p2 = tf1.add_paragraph()
p2.text = "The Talent Governance & 1:1 Execution Platform"
p2.font.size = Pt(22)
p2.font.bold = True
p2.font.color.rgb = RGBColor(129, 140, 248)
p2.space_before = Pt(8)

p3 = tf1.add_paragraph()
p3.text = "Bridging the gap between HR systems of record and ground manager follow-through — driving continuous feedback, living 1:1 agendas, and evidence-based talent evaluations."
p3.font.size = Pt(13)
p3.font.color.rgb = RGBColor(203, 213, 225)
p3.space_before = Pt(14)

p4 = tf1.add_paragraph()
p4.text = "Continuous 1:1s  •  Named Feedback  •  Micro-Win Logs  •  PIP Governance  •  Training Demand Matrix"
p4.font.size = Pt(11)
p4.font.bold = True
p4.font.color.rgb = RGBColor(148, 163, 184)
p4.space_before = Pt(32)

# ==============================================================================
# SLIDE 2: The Core HR Execution Gap
# ==============================================================================
s2 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s2, COLOR_BG_LIGHT)
add_header(s2, "01 — The HR Execution Gap", "HR Has the Systems of Record. The Breakdown Is Manager Follow-Through.")

c2_left = add_card(s2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
tf2_l = c2_left.text_frame
tf2_l.word_wrap = True
tf2_l.margin_left = Inches(0.4)
tf2_l.margin_top = Inches(0.35)

p = tf2_l.paragraphs[0]
p.text = "Why Traditional HR Tech Fails Line Teams:"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY

pains = [
    ("Scattered Meeting Notes", "1:1 notes live in private Google Docs, paper diaries, or forgotten Slack messages with zero task follow-through."),
    ("Recency Bias in Appraisals", "Annual reviews only remember the last 3 weeks; consistent contributions from earlier in the year are forgotten."),
    ("The 'Chasing Bot' HR Problem", "HR spends 60% of their time chasing line managers for review sign-offs, probation checks, and PIP updates."),
    ("Toxic Anonymous Guessing", "Anonymous feedback creates paranoia; teams need constructive, transparent feedback tied to real accountability.")
]
for title, desc in pains:
    pt = tf2_l.add_paragraph()
    pt.text = f"• {title}: "
    pt.font.size = Pt(11)
    pt.font.bold = True
    pt.font.color.rgb = COLOR_ACCENT
    pt.space_before = Pt(8)
    
    pd = tf2_l.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(10)
    pd.font.color.rgb = COLOR_MUTED

c2_right = add_card(s2, Inches(6.7), Inches(1.8), Inches(5.833), Inches(4.8), RGBColor(238, 242, 255), RGBColor(199, 210, 254))
tf2_r = c2_right.text_frame
tf2_r.word_wrap = True
tf2_r.margin_left = Inches(0.4)
tf2_r.margin_top = Inches(0.4)

p = tf2_r.paragraphs[0]
p.text = "THE DAIKO SOLUTION"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT

p_sol = tf2_r.add_paragraph()
p_sol.text = "A Purpose-Built Execution Layer Above Your HRMS"
p_sol.font.size = Pt(18)
p_sol.font.bold = True
p_sol.font.color.rgb = COLOR_PRIMARY
p_sol.space_before = Pt(8)

p_body = tf2_r.add_paragraph()
p_body.text = "DAIKO does not replace Darwinbox, Workday, or SAP. It gives managers and employees a lightweight shared space to run continuous 1:1s, capture real-time wins, and ensure every review is backed by factual operational data."
p_body.font.size = Pt(12)
p_body.font.color.rgb = COLOR_PRIMARY
p_body.space_before = Pt(12)

# ==============================================================================
# SLIDE 3: Role Separation & Shared Harmony
# ==============================================================================
s3 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s3, COLOR_BG_LIGHT)
add_header(s3, "02 — Role Harmony", "Two Dedicated Views. One Unified Performance Engine.")

# Manager Card
c3_mgr = add_card(s3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
tf3_m = c3_mgr.text_frame
tf3_m.word_wrap = True
tf3_m.margin_left = Inches(0.35)
tf3_m.margin_top = Inches(0.3)

p = tf3_m.paragraphs[0]
p.text = "👔 MANAGER HUB (Team Orchestration)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT

mgr_features = [
    "• Team 1:1 Cadence RAG (Flags >3 weeks of missed syncs)",
    "• Living 1:1 Agenda & Action Follow-Through Tracker",
    "• Named Feedback & Kudos Stream (Give & Review)",
    "• Micro-Win Log (1-click capture of employee achievements)",
    "• PIP Milestone & Check-in Tracker",
    "• Skill Gap & Training Need Tagging"
]
for item in mgr_features:
    pi = tf3_m.add_paragraph()
    pi.text = item
    pi.font.size = Pt(11)
    pi.font.color.rgb = COLOR_PRIMARY
    pi.space_before = Pt(8)

# Employee Card
c3_emp = add_card(s3, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.8))
tf3_e = c3_emp.text_frame
tf3_e.word_wrap = True
tf3_e.margin_left = Inches(0.35)
tf3_e.margin_top = Inches(0.3)

p = tf3_e.paragraphs[0]
p.text = "🌱 EMPLOYEE PORTAL (Growth & Clarity)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = COLOR_TEAL

emp_features = [
    "• Upcoming 1:1 Agenda Scratchpad (Add topics anytime)",
    "• My Commitments Checklist (Action items from last 1:1)",
    "• Named Kudos & Recognition Received (Permanent profile record)",
    "• Personal Skill Wishlist & Training Catalog",
    "• Objective Milestone Progress (Goal / PIP Tracking)",
    "• Continuous Self-Appraisal Data Bank"
]
for item in emp_features:
    pi = tf3_e.add_paragraph()
    pi.text = item
    pi.font.size = Pt(11)
    pi.font.color.rgb = COLOR_PRIMARY
    pi.space_before = Pt(8)

# ==============================================================================
# SLIDE 4: Differentiating 1:1 Notes vs. Continuous Feedback
# ==============================================================================
s4 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s4, COLOR_BG_LIGHT)
add_header(s4, "03 — Process Clarity", "Meeting Notes ≠ Feedback: Keeping Both Clean and Distinct")

c4_notes = add_card(s4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
tf4_n = c4_notes.text_frame
tf4_n.word_wrap = True
tf4_n.margin_left = Inches(0.35)
tf4_n.margin_top = Inches(0.3)

p = tf4_n.paragraphs[0]
p.text = "📝 1:1 MEETING NOTES & ACTIONS"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT

notes_points = [
    ("What it is", "Operational commitments, blockers, and project deliverables."),
    ("Privacy Level", "Strictly private between Manager & Employee."),
    ("Format", "Living agenda bullet points that convert into tracked action items."),
    ("Carryover Engine", "Unfinished tasks auto-roll to next week's sync so nothing is lost."),
    ("HR Visibility", "HR tracks task closure % and meeting cadence, NOT private note text.")
]
for tag, desc in notes_points:
    pt = tf4_n.add_paragraph()
    pt.text = f"{tag}: "
    pt.font.size = Pt(11)
    pt.font.bold = True
    pt.font.color.rgb = COLOR_PRIMARY
    pt.space_before = Pt(6)
    
    pd = tf4_n.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(10)
    pd.font.color.rgb = COLOR_MUTED

c4_feed = add_card(s4, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.8))
tf4_f = c4_feed.text_frame
tf4_f.word_wrap = True
tf4_f.margin_left = Inches(0.35)
tf4_f.margin_top = Inches(0.3)

p = tf4_f.paragraphs[0]
p.text = "💬 CONTINUOUS NAMED FEEDBACK"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = COLOR_TEAL

feed_points = [
    ("What it is", "Behavioral observations, praise for good work, or coaching nudges."),
    ("Transparency", "Always attached with real author names (100% accountable)."),
    ("Format", "Bite-sized 30-word cards using the Situation-Behavior-Impact model."),
    ("Where it lives", "Logs permanently to the employee's 360 Profile and Annual Appraisal."),
    ("HR Visibility", "Rolls up into overall team sentiment and culture health metrics.")
]
for tag, desc in feed_points:
    pt = tf4_f.add_paragraph()
    pt.text = f"{tag}: "
    pt.font.size = Pt(11)
    pt.font.bold = True
    pt.font.color.rgb = COLOR_PRIMARY
    pt.space_before = Pt(6)
    
    pd = tf4_f.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(10)
    pd.font.color.rgb = COLOR_MUTED

# ==============================================================================
# SLIDE 5: The Micro-Win Log & Fairness Safeguard
# ==============================================================================
s5 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s5, COLOR_BG_LIGHT)
add_header(s5, "04 — Objectivity & Fairness", "The Micro-Win Log: Safeguarding Great Work from Recency Bias")

cards5_data = [
    ("THE RECENCY BIAS TRAP", "Managers often rate employees on their last 3 weeks of performance, completely forgetting major wins from months prior. Struggling employees are often unfairly stereotyped as having done 'nothing good.'", COLOR_AMBER),
    ("THE 1-CLICK 'WIN NOTE'", "Whenever an employee resolves a crisis, assists a teammate, or delivers client praise, the manager logs a 10-second Win Note with tags (e.g. #HighOwnership, #QuickDelivery).", COLOR_ACCENT),
    ("EVIDENCE-BASED REVIEWS", "During annual appraisals or PIP reviews, DAIKO generates a chronological timeline of all verified wins alongside development areas—ensuring 100% objective, defensible talent decisions.", COLOR_GREEN)
]

for i, (title, desc, color) in enumerate(cards5_data):
    c = add_card(s5, Inches(0.8 + i * 4.0), Inches(1.9), Inches(3.733), Inches(3.6))
    tf = c.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.3)
    
    p = tf.paragraphs[0]
    p.text = f"STEP 0{i+1}"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = color
    
    pt = tf.add_paragraph()
    pt.text = title
    pt.font.size = Pt(14)
    pt.font.bold = True
    pt.font.color.rgb = COLOR_PRIMARY
    pt.space_before = Pt(6)
    
    pd = tf.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(11)
    pd.font.color.rgb = COLOR_MUTED
    pd.space_before = Pt(10)

callout5 = add_card(s5, Inches(0.8), Inches(5.8), Inches(11.733), Inches(0.9), RGBColor(241, 245, 249))
tf_c5 = callout5.text_frame
tf_c5.word_wrap = True
p = tf_c5.paragraphs[0]
p.text = "Fairness Guarantee: Every employee—even one on a PIP—has their positive contributions documented and accounted for."
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY
p.alignment = PP_ALIGN.CENTER

# ==============================================================================
# SLIDE 6: 1:1 Execution Statistics (What HR & Leadership Sees)
# ==============================================================================
s6 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s6, COLOR_BG_LIGHT)
add_header(s6, "05 — 1:1 Health & Execution", "1:1 Follow-Through Statistics: Measuring What Was Once Invisible")

stats_data = [
    ("91.4%", "1:1 Cadence Adherence", "Target: >85% completed bi-weekly syncs across all teams."),
    ("86.2%", "Action Item Closure Rate", "Agreed 1:1 tasks marked completed within target deadline."),
    ("14.5 Days", "Avg. Time Between 1:1s", "Early alert triggers if a manager exceeds 21 days without a sync."),
    ("100%", "Private Note Integrity", "Zero note text visible to HR; metrics track operational habits only.")
]

for i, (metric, label, sub) in enumerate(stats_data):
    c = add_card(s6, Inches(0.8 + i * 2.98), Inches(1.9), Inches(2.78), Inches(3.6))
    tf = c.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.3)
    
    p = tf.paragraphs[0]
    p.text = metric
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT
    
    pt = tf.add_paragraph()
    pt.text = label
    pt.font.size = Pt(13)
    pt.font.bold = True
    pt.font.color.rgb = COLOR_PRIMARY
    pt.space_before = Pt(8)
    
    pd = tf.add_paragraph()
    pd.text = sub
    pd.font.size = Pt(10)
    pd.font.color.rgb = COLOR_MUTED
    pd.space_before = Pt(8)

callout6 = add_card(s6, Inches(0.8), Inches(5.8), Inches(11.733), Inches(0.9), RGBColor(238, 242, 255), RGBColor(199, 210, 254))
tf_c6 = callout6.text_frame
tf_c6.word_wrap = True
p = tf_c6.paragraphs[0]
p.text = "Automated Calendar Sync: Detects scheduled 1:1 calls on Google Meet / Teams and nudges both parties for instant follow-up."
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT
p.alignment = PP_ALIGN.CENTER

# ==============================================================================
# SLIDE 7: Talent Profile 360, PIP & Appraisal Suite
# ==============================================================================
s7 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s7, COLOR_BG_LIGHT)
add_header(s7, "06 — Performance Governance", "Talent Profiles, Structured PIPs & Fact-Based Appraisals")

cards7_cols = [
    ("TALENT PROFILE 360", [
        "• Unified view of verified skills, certifications & project history",
        "• Chronological timeline of peer kudos & manager win notes",
        "• Historical 1:1 engagement and task closure score",
        "• Readiness rating for promotion and cross-functional mobility"
    ]),
    ("STRUCTURED PIP TRACKER", [
        "• Clear, time-bound milestones (30/60/90 days)",
        "• Weekly check-in agendas with objective completion meters",
        "• Eliminates ambiguity—both manager & employee see exact goals",
        "• Formal sign-off workflow with HR governance oversight"
    ]),
    ("APPRAISAL CALIBRATION", [
        "• Auto-aggregates 12 months of 1:1 goals, tasks, and win logs",
        "• Removes 90% of annual review preparation friction",
        "• Prevents favoritism or unfair ratings with auditable proof",
        "• 1-tap sync back into your primary HRMS of record"
    ])
]

for i, (title, items) in enumerate(cards7_cols):
    c = add_card(s7, Inches(0.8 + i * 4.0), Inches(1.8), Inches(3.733), Inches(4.8))
    tf = c.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.35)
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT
    
    for item in items:
        pi = tf.add_paragraph()
        pi.text = item
        pi.font.size = Pt(11)
        pi.font.color.rgb = COLOR_PRIMARY
        pi.space_before = Pt(8)

# ==============================================================================
# SLIDE 8: Training Needs Identification (TNI)
# ==============================================================================
s8 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s8, COLOR_BG_LIGHT)
add_header(s8, "07 — Skill Intelligence", "Training Needs Identification: Bottom-Up Demand Aggregation")

c8_left = add_card(s8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
tf8_l = c8_left.text_frame
tf8_l.word_wrap = True
tf8_l.margin_left = Inches(0.35)
tf8_l.margin_top = Inches(0.3)

p = tf8_l.paragraphs[0]
p.text = "How Training Demand Is Captured:"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY

tni_steps = [
    ("1. Spot During 1:1s", "When a manager or employee identifies a skill hurdle during a project review, they tag a 'Skill Need' in 1 click."),
    ("2. Employee Wishlists", "Team members can self-select training modules or domain certifications they want to pursue for career growth."),
    ("3. Auto-Aggregation", "DAIKO clusters individual skill tags across departments to show HR exactly which training cohorts to organize.")
]
for title, desc in tni_steps:
    pt = tf8_l.add_paragraph()
    pt.text = title
    pt.font.size = Pt(12)
    pt.font.bold = True
    pt.font.color.rgb = COLOR_ACCENT
    pt.space_before = Pt(8)
    
    pd = tf8_l.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(10)
    pd.font.color.rgb = COLOR_MUTED

c8_right = add_card(s8, Inches(6.7), Inches(1.8), Inches(5.833), Inches(4.8), RGBColor(241, 245, 249))
tf8_r = c8_right.text_frame
tf8_r.word_wrap = True
tf8_r.margin_left = Inches(0.35)
tf8_r.margin_top = Inches(0.3)

p = tf8_r.paragraphs[0]
p.text = "LIVE TRAINING DEMAND MATRIX (HR VIEW)"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_MUTED

matrix_items = [
    "• Cloud Architecture & Security: 18 Engineers (Across 4 Teams)",
    "• Client De-escalation & Communication: 12 PMs & Leads",
    "• Advanced Financial Modeling: 7 Analysts",
    "• People Leadership & Feedback Coaching: 9 New Managers"
]
for item in matrix_items:
    pi = tf8_r.add_paragraph()
    pi.text = item
    pi.font.size = Pt(11)
    pi.font.bold = True
    pi.font.color.rgb = COLOR_PRIMARY
    pi.space_before = Pt(10)

p_btn = tf8_r.add_paragraph()
p_btn.text = "HR Action: [ Schedule Cohort Workshop ]  [ Assign Budget ]"
p_btn.font.size = Pt(11)
p_btn.font.bold = True
p_btn.font.color.rgb = COLOR_ACCENT
p_btn.space_before = Pt(18)

# ==============================================================================
# SLIDE 9: Talent Insights & Governance Dashboard
# ==============================================================================
s9 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s9, COLOR_BG_LIGHT)
add_header(s9, "08 — Executive Visibility", "Talent Insights & Governance: Real-Time People Intelligence")

dash_cols = [
    ("NOW (What's happening?)", [
        "• 4 Managers have not held a 1:1 in >21 days",
        "• 12 Micro-Wins logged across Engineering today",
        "• 3 Active PIPs approaching 30-day milestone review"
    ]),
    ("WHY (What's causing it?)", [
        "• Team A onboarding load causing sprint schedule drift",
        "• Q3 release milestone driving high employee overtime",
        "• Skill gap identified in frontend architecture testing"
    ]),
    ("NEXT (What needs action?)", [
        "• Automated nudge sent to Managers to schedule 1:1s",
        "• 1-Tap Director approval for high-performer spot bonus",
        "• Batch 18 engineers for Cloud Architecture training"
    ])
]

for i, (title, items) in enumerate(dash_cols):
    c = add_card(s9, Inches(0.8 + i * 4.0), Inches(1.8), Inches(3.733), Inches(3.6))
    tf = c.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.3)
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT
    
    for item in items:
        pi = tf.add_paragraph()
        pi.text = item
        pi.font.size = Pt(11)
        pi.font.color.rgb = COLOR_PRIMARY
        pi.space_before = Pt(8)

callout9 = add_card(s9, Inches(0.8), Inches(5.8), Inches(11.733), Inches(0.9), RGBColor(15, 23, 42), RGBColor(30, 41, 59))
tf_c9 = callout9.text_frame
tf_c9.word_wrap = True
p = tf_c9.paragraphs[0]
p.text = "ORGANIZATION VITALS:   1:1 Cadence: 91.4%   |   Action Closure: 86.2%   |   Active PIPs: 03   |   Pending Appraisals: 00"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = RGBColor(129, 140, 248)
p.alignment = PP_ALIGN.CENTER

# ==============================================================================
# SLIDE 10: The 30-Day Proof of Value (The Close)
# ==============================================================================
s10 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s10, COLOR_BG_DARK)
add_header(s10, "09 — Implementation Blueprint", "30-Day Proof of Value", dark=True)

pov_items = [
    ("ONE DEPARTMENT", "Roll out across 1 business unit or engineering department."),
    ("ONE WORKFLOW", "Living 1:1 cadence, action tracking & continuous micro-wins."),
    ("ONE BASELINE", "Historical 1:1 frequency and task completion benchmarks."),
    ("ONE SCORECARD", "1:1 Cadence %, Action Closure %, and Manager Adoption Rate."),
    ("ONE DECISION", "Scale across company / Refine workflows / Stop.")
]

for i, (title, desc) in enumerate(pov_items):
    c = add_card(s10, Inches(0.8 + i * 2.38), Inches(1.9), Inches(2.18), Inches(3.4), RGBColor(30, 41, 59), RGBColor(51, 65, 85))
    tf = c.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.3)
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(129, 140, 248)
    
    pd = tf.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(11)
    pd.font.color.rgb = RGBColor(203, 213, 225)
    pd.space_before = Pt(8)

callout10 = add_card(s10, Inches(0.8), Inches(5.6), Inches(11.733), Inches(1.1), RGBColor(30, 41, 59), RGBColor(129, 140, 248))
tf_c10 = callout10.text_frame
tf_c10.word_wrap = True
p = tf_c10.paragraphs[0]
p.text = "The Executive Guarantee:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = RGBColor(129, 140, 248)
p.alignment = PP_ALIGN.CENTER

p2 = tf_c10.add_paragraph()
p2.text = "\"No HRMS replacement. Zero operational disruption. Measurable manager follow-through before scale.\""
p2.font.size = Pt(14)
p2.font.bold = True
p2.font.color.rgb = RGBColor(255, 255, 255)
p2.alignment = PP_ALIGN.CENTER

# 3. Save Output Presentation
output_filename = "DAIKO_People_Talent_Governance_Deck.pptx"
prs.save(output_filename)
print(f"Presentation saved successfully as '{output_filename}'")